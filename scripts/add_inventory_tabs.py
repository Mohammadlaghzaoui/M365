# -*- coding: utf-8 -*-
"""
Adds admin-center inventory tabs (Exchange, SharePoint, Teams, Azure) to the
M365 migration workbooks. Existing sheets and data are never modified or
removed — tabs are only added (skipped if already present).
"""
import os
import sys
import glob
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

C_DARK = "1F3864"
C_MED = "2F5597"

F_TITLE = Font(name="Calibri", size=20, bold=True, color="FFFFFF")
F_SUB = Font(name="Calibri", size=12, bold=False, color="FFFFFF")
F_HDR = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
F_BODY = Font(name="Calibri", size=11)

FILL_DARK = PatternFill("solid", fgColor=C_DARK)
FILL_MED = PatternFill("solid", fgColor=C_MED)

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

DATA_ROWS = 150


def banner(ws, company, subtitle, ncols):
    last = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    c = ws["A1"]
    c.value = f"M365 Tenant-to-Tenant Migration  |  {company}  →  Kelso"
    c.font = F_TITLE
    c.fill = FILL_DARK
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    s = ws["A2"]
    s.value = subtitle
    s.font = F_SUB
    s.fill = FILL_MED
    s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 20
    for col in range(1, ncols + 1):
        ws.cell(row=1, column=col).fill = FILL_DARK
        ws.cell(row=2, column=col).fill = FILL_MED


def build_inventory_sheet(wb, index, company, title, tab_color, subtitle,
                          headers, widths, dropdowns):
    """dropdowns: {column_letter: 'comma,separated,values'}"""
    ws = wb.create_sheet(title, index)
    ws.sheet_properties.tabColor = tab_color
    banner(ws, company, subtitle, len(headers))
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=4, column=i, value=h)
        cell.font = F_HDR
        cell.fill = FILL_MED
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[4].height = 28
    first, last = 5, 4 + DATA_ROWS
    for r in range(first, last + 1):
        ws.cell(row=r, column=1, value=r - 4).alignment = CENTER
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=col)
            cell.border = BORDER
            cell.font = F_BODY
    for col_letter, values in dropdowns.items():
        dv = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}{first}:{col_letter}{last}")
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{last}"
    return ws


SHEETS = [
    ("Exchange Inventory", "ED7D31",
     "Exchange admin center — mailbox inventory for analysis (paste exports here)",
     ["No", "Display name", "Primary SMTP address", "Mailbox type", "Aliases",
      "Mailbox size (GB)", "Item count", "Archive (GB)", "Litigation hold",
      "Forwarding to", "Full Access delegates", "Send As", "Last logon",
      "Migrate?", "Notes"],
     [6, 24, 34, 14, 28, 14, 12, 12, 12, 26, 26, 22, 14, 11, 28],
     {"D": "User,Shared,Room,Equipment",
      "I": "Yes,No",
      "N": "Yes,No,In progress"}),
    ("SharePoint Inventory", "5B9BD5",
     "SharePoint admin center — site inventory for analysis (paste exports here)",
     ["No", "Site name", "Site URL", "Type", "Size (GB)", "Files",
      "Owner(s)", "External sharing", "Last activity", "Connected to Team",
      "Unique permissions", "Migrate?", "Target site in Kelso", "Notes"],
     [6, 26, 42, 18, 12, 12, 26, 15, 14, 14, 14, 11, 34, 28],
     {"D": "Team site,Communication site,Classic site,Hub site,Other",
      "H": "On,Off",
      "J": "Yes,No",
      "K": "Yes,No",
      "L": "Yes,No,Archive"}),
    ("Teams Inventory", "7030A0",
     "Teams admin center — team inventory for analysis (paste exports here)",
     ["No", "Team name", "Privacy", "Owners", "Members", "Guests",
      "Channels", "Private/shared channels", "SharePoint URL", "Tabs & apps",
      "Chat history in scope", "Migrate?", "Notes"],
     [6, 28, 11, 26, 10, 10, 10, 16, 42, 26, 14, 11, 28],
     {"C": "Public,Private",
      "H": "Yes,No",
      "K": "Yes,No",
      "L": "Yes,No"}),
    ("Azure Inventory", "0078D4",
     "Azure / Entra ID — identity & platform inventory for analysis (paste exports here)",
     ["No", "Category", "Name", "Details / configuration", "Owner",
      "Members / assigned users", "Status in source", "Required in Kelso?",
      "Action", "Notes"],
     [6, 22, 30, 44, 22, 26, 16, 14, 12, 28],
     {"B": ("Security group,M365 group,Distribution list,Enterprise app,"
            "App registration,Conditional Access policy,Admin role,Device,"
            "Domain,Intune policy,License,Other"),
      "H": "Yes,No,TBD",
      "I": "Rebuild,Migrate,Skip"}),
]


def add_tabs(path, out_path=None):
    wb = load_workbook(path)
    company = None
    if "Overview" in wb.sheetnames:
        for row in wb["Overview"].iter_rows(min_row=4, max_row=20):
            if row[0].value == "Source tenant":
                company = row[1].value
                break
    if not company:
        company = os.path.basename(path).replace("M365_Migration_", "").split("_to_")[0].replace("_", " ")

    anchor = wb.sheetnames.index("User Mapping") if "User Mapping" in wb.sheetnames else len(wb.sheetnames) - 1
    added = []
    offset = 1
    for title, color, subtitle, headers, widths, dvs in SHEETS:
        if title in wb.sheetnames:
            continue
        build_inventory_sheet(wb, anchor + offset, company, title, color,
                              subtitle, headers, widths, dvs)
        added.append(title)
        offset += 1
    wb.save(out_path or path)
    print(f"{os.path.basename(out_path or path)}: added {added if added else 'nothing (already present)'}")


if __name__ == "__main__":
    if len(sys.argv) > 2:
        add_tabs(sys.argv[1], sys.argv[2])
    elif len(sys.argv) > 1:
        add_tabs(sys.argv[1])
    else:
        base = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "Migration-Excels")
        for f in sorted(glob.glob(os.path.join(base, "*.xlsx"))):
            add_tabs(f)
