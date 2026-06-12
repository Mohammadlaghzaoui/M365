# -*- coding: utf-8 -*-
"""
Genereert per bedrijf een professioneel M365 tenant-naar-tenant migratiewerkboek.
Bron-tenant: het bedrijf zelf | Doel-tenant: Kelso
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

COMPANIES = [
    "RGD", "Retrofit", "Quality H&A", "RMR", "LLR", "YES", "Frey Lutz",
    "Icon", "Wallace", "Markade", "Fab-Logix", "Premier", "Strictly",
]
TARGET = "Kelso"
OUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "Migratie-Excels")

# ---------- Stijl-definities ----------
C_DARK = "1F3864"      # donkerblauw
C_MED = "2F5597"       # middenblauw
C_LIGHT = "D6E4F0"     # lichtblauw
C_ACCENT = "ED7D31"    # oranje accent
C_GREY = "F2F2F2"

F_TITLE = Font(name="Calibri", size=20, bold=True, color="FFFFFF")
F_SUB = Font(name="Calibri", size=12, bold=False, color="FFFFFF")
F_HDR = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
F_PHASE = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
F_BODY = Font(name="Calibri", size=11)
F_BOLD = Font(name="Calibri", size=11, bold=True)

FILL_DARK = PatternFill("solid", fgColor=C_DARK)
FILL_MED = PatternFill("solid", fgColor=C_MED)
FILL_LIGHT = PatternFill("solid", fgColor=C_LIGHT)
FILL_ACCENT = PatternFill("solid", fgColor=C_ACCENT)
FILL_GREY = PatternFill("solid", fgColor=C_GREY)

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

STATUS_LIST = '"Niet gestart,Bezig,Voltooid,Geblokkeerd,N.v.t."'
PRIO_LIST = '"Hoog,Middel,Laag"'

STATUS_COLORS = {
    "Voltooid": "C6EFCE",
    "Bezig": "FFEB9C",
    "Geblokkeerd": "FFC7CE",
}


def status_validation(ws, col, first_row, last_row):
    dv = DataValidation(type="list", formula1=STATUS_LIST, allow_blank=True)
    dv.error = "Kies een waarde uit de lijst"
    dv.errorTitle = "Ongeldige status"
    ws.add_data_validation(dv)
    dv.add(f"{col}{first_row}:{col}{last_row}")
    for value, color in STATUS_COLORS.items():
        ws.conditional_formatting.add(
            f"{col}{first_row}:{col}{last_row}",
            CellIsRule(operator="equal", formula=[f'"{value}"'],
                       fill=PatternFill("solid", fgColor=color)))


def prio_validation(ws, col, first_row, last_row):
    dv = DataValidation(type="list", formula1=PRIO_LIST, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}{first_row}:{col}{last_row}")
    ws.conditional_formatting.add(
        f"{col}{first_row}:{col}{last_row}",
        CellIsRule(operator="equal", formula=['"Hoog"'],
                   fill=PatternFill("solid", fgColor="FFC7CE")))


def banner(ws, company, subtitle, ncols):
    last = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    c = ws["A1"]
    c.value = f"M365 Tenant-naar-Tenant Migratie  |  {company}  →  {TARGET}"
    c.font = F_TITLE
    c.fill = FILL_DARK
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    s = ws["A2"]
    s.value = subtitle
    s.font = F_SUB
    s.fill = FILL_MED
    s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 20
    for col in range(1, ncols + 1):
        ws.cell(row=1, column=col).fill = FILL_DARK
        ws.cell(row=2, column=col).fill = FILL_MED


def header_row(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=row, column=i, value=h)
        cell.font = F_HDR
        cell.fill = FILL_MED
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 22


def task_sheet(wb, company, title, tab_color, intro, phases):
    """phases: list of (fase-naam, [(taak, beschrijving, verantwoordelijke-default), ...])"""
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab_color
    headers = ["Nr", "Taak", "Beschrijving / Instructie", "Verantwoordelijke",
               "Prioriteit", "Status", "Deadline", "Afgerond op", "Notities"]
    widths = [6, 38, 70, 20, 11, 13, 13, 13, 35]
    banner(ws, company, intro, len(headers))
    header_row(ws, 4, headers, widths)

    row = 5
    nr = 1
    for phase_name, tasks in phases:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        pc = ws.cell(row=row, column=1, value=phase_name)
        pc.font = F_PHASE
        pc.fill = FILL_ACCENT
        pc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in range(1, len(headers) + 1):
            ws.cell(row=row, column=col).fill = FILL_ACCENT
        ws.row_dimensions[row].height = 20
        row += 1
        for task, desc, owner in tasks:
            values = [nr, task, desc, owner, "", "Niet gestart", "", "", ""]
            for col, v in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=v)
                cell.font = F_BODY
                cell.border = BORDER
                cell.alignment = WRAP if col in (2, 3, 9) else CENTER
            if row % 2 == 0:
                for col in range(1, len(headers) + 1):
                    if ws.cell(row=row, column=col).fill.fgColor.rgb in (None, "00000000"):
                        ws.cell(row=row, column=col).fill = FILL_GREY
            nr += 1
            row += 1

    last_row = row - 1
    status_validation(ws, "F", 5, last_row)
    prio_validation(ws, "E", 5, last_row)
    ws.auto_filter.ref = f"A4:I{last_row}"
    ws.freeze_panes = "A5"
    return ws


# ---------- Inhoud van de fases ----------

def phases_voorbereiding(company):
    return [
        ("FASE 1 — PROJECT & GOVERNANCE", [
            ("Kick-off & scope vastleggen",
             f"Projectscope bepalen: welke workloads van {company} migreren naar {TARGET} (Exchange, OneDrive, SharePoint, Teams, apps). Stakeholders en beslissers benoemen.",
             "Projectleider"),
            ("Migratieplanning & tijdlijn",
             "Gedetailleerde planning opstellen met mijlpalen, migratiegolven (waves) en cutover-datum. Rekening houden met vakanties en bedrijfskritische periodes.",
             "Projectleider"),
            ("Communicatieplan opstellen",
             "Eindgebruikers tijdig informeren: wat verandert er, wanneer, wat moeten ze zelf doen (nieuw wachtwoord, MFA opnieuw instellen, Outlook-profiel, Teams).",
             "Projectleider"),
            ("Rollback-/noodplan definiëren",
             "Terugvalscenario beschrijven per workload als de migratie misloopt (MX terugzetten, forwarding, herstel vanaf back-up).",
             "M365 Engineer"),
            ("Migratietool kiezen & licenties",
             "Tool selecteren en licenties aanschaffen (bv. BitTitan MigrationWiz, Quest On Demand, AvePoint Fly, ShareGate). Aantal mailboxen/users bepalen voor licentietelling.",
             "M365 Engineer"),
        ]),
        ("FASE 2 — TOEGANG & ACCOUNTS", [
            ("Global Admin toegang bron-tenant",
             f"Globale beheerder-account in de {company}-tenant verkrijgen/aanmaken (incl. MFA). Break-glass account documenteren.",
             "M365 Engineer"),
            (f"Global Admin toegang doel-tenant ({TARGET})",
             f"Globale beheerder-account in de {TARGET}-tenant. Controleren dat er voldoende rechten zijn voor user-creatie, Exchange, SharePoint en Teams admin.",
             "M365 Engineer"),
            ("Service-accounts voor migratietool",
             "Dedicated migratie-admin accounts aanmaken in bron én doel (zonder MFA-blokkade voor de tool, of met app-registratie/modern auth). Na afloop verwijderen.",
             "M365 Engineer"),
            ("App-registraties & API-permissies",
             "App-registratie in Entra ID voor de migratietool met de juiste Graph/EWS-permissies; admin consent geven in beide tenants.",
             "M365 Engineer"),
        ]),
        ("FASE 3 — LICENTIES & CAPACITEIT DOEL-TENANT", [
            (f"Licenties controleren in {TARGET}",
             f"Voldoende M365-licenties (Business Premium/E3/E5 etc.) beschikbaar in {TARGET} voor alle te migreren gebruikers van {company}.",
             "Projectleider"),
            ("Opslagcapaciteit SharePoint controleren",
             f"Controle of de SharePoint-opslagquota in {TARGET} de data van {company} aankunnen (inventory-cijfers gebruiken).",
             "M365 Engineer"),
            ("Naamgevingsconventie bepalen",
             f"UPN-/e-mailformaat in {TARGET} vastleggen (bv. voornaam.achternaam@kelso-domein) en conflicten met bestaande {TARGET}-gebruikers identificeren.",
             "M365 Engineer"),
        ]),
        ("FASE 4 — DOMEIN & DNS VOORBEREIDING", [
            ("DNS-beheer toegang regelen",
             f"Toegang tot de DNS-registrar/zone van het {company}-domein verifiëren. TTL van MX/Autodiscover records vóór cutover verlagen naar 300-3600 sec.",
             "M365 Engineer"),
            ("Domeinverhuizing voorbereiden",
             f"Plan voor het verwijderen van het maildomein uit de {company}-tenant en toevoegen + verifiëren in {TARGET} (verificatie-TXT klaarzetten). Let op: domein kan pas verhuizen als alle objecten het niet meer gebruiken.",
             "M365 Engineer"),
            ("Externe afhankelijkheden DNS",
             "SPF, DKIM, DMARC, MTA-STS en eventuele 3rd-party records (printers/scanners, applicaties die SMTP-relay gebruiken) inventariseren en herconfigureren.",
             "M365 Engineer"),
        ]),
        ("FASE 5 — EINDGEBRUIKERS & WERKPLEK", [
            ("Workstations inventariseren",
             "Overzicht van alle pc's/laptops: Entra-joined of hybrid? Intune-managed? Plan voor het omhangen van devices naar de Kelso-tenant.",
             "Werkplekbeheer"),
            ("Outlook-/OneDrive-profielen plan",
             "Procedure opstellen voor nieuw Outlook-profiel en OneDrive-herkoppeling op elke werkplek na cutover (handmatig of gescript).",
             "Werkplekbeheer"),
            ("Mobiele apparaten",
             "Plan voor het opnieuw inschrijven van telefoons/tablets (Outlook mobile, Teams, Company Portal, MFA/Authenticator opnieuw registreren).",
             "Werkplekbeheer"),
            ("Wachtwoorden & MFA-instructies",
             "Instructiedocument voor gebruikers: eerste aanmelding in Kelso-tenant, wachtwoord instellen, MFA registreren.",
             "Servicedesk"),
        ]),
    ]


def phases_inventory(company):
    return [
        ("INVENTORY — IDENTITEIT", [
            ("Gebruikersoverzicht exporteren",
             f"Alle gebruikers uit {company} exporteren (Entra ID > Users of Graph/PowerShell: Get-MgUser). Vastleggen: UPN, displaynaam, licenties, laatste login, actief/inactief.",
             "M365 Engineer"),
            ("Gedeelde mailboxen & resources",
             "Shared mailboxes, room/equipment mailboxes en hun delegaties (Full Access, Send As, Send on Behalf) exporteren (Get-Mailbox / Get-MailboxPermission).",
             "M365 Engineer"),
            ("Distributielijsten & M365-groepen",
             "Alle distributiegroepen, mail-enabled security groups en Microsoft 365-groepen incl. leden en eigenaren exporteren.",
             "M365 Engineer"),
            ("Gast-accounts (B2B)",
             "Externe gasten in de tenant inventariseren; bepalen welke opnieuw uitgenodigd moeten worden in Kelso.",
             "M365 Engineer"),
            ("Admin-rollen documenteren",
             "Wie heeft welke beheerrol in de bron-tenant; bepalen welke rollen in Kelso nodig zijn.",
             "M365 Engineer"),
        ]),
        ("INVENTORY — EXCHANGE ONLINE", [
            ("Mailboxgroottes & itemcounts",
             "Per mailbox grootte en aantal items exporteren (Get-MailboxStatistics) — bepaalt migratieduur en eventuele pre-staging.",
             "M365 Engineer"),
            ("Archiefmailboxen",
             "In-Place Archives identificeren (grootte, auto-expanding archives vergen extra aandacht/tooling).",
             "M365 Engineer"),
            ("Mailflow-regels & connectors",
             "Transport rules, connectors, anti-spam/anti-phish policies en e-maildisclaimers documenteren om na te bouwen in Kelso.",
             "M365 Engineer"),
            ("Forwarding & inbox rules",
             "Mailbox-forwarding en kritische inbox rules inventariseren (gaan niet altijd automatisch mee).",
             "M365 Engineer"),
            ("Aliassen & proxy-adressen",
             "Alle SMTP-aliassen per mailbox exporteren zodat ze in Kelso opnieuw aangemaakt kunnen worden.",
             "M365 Engineer"),
            ("Litigation hold & retentie",
             "Mailboxen met Litigation Hold/retention policies identificeren — data veiligstellen vóór migratie (hold-data migreert niet vanzelf).",
             "M365 Engineer"),
        ]),
        ("INVENTORY — ONEDRIVE & SHAREPOINT", [
            ("OneDrive-overzicht",
             "Per gebruiker OneDrive-URL, grootte en aantal bestanden exporteren (SharePoint admin center of PowerShell).",
             "M365 Engineer"),
            ("SharePoint-sites inventariseren",
             "Alle sites (team/communicatie), grootte, eigenaren, laatste activiteit. Verouderde sites markeren om NIET te migreren (opschonen = sneller migreren).",
             "M365 Engineer"),
            ("Permissies & externe deellinks",
             "Unieke permissies, externe sharing en anonieme links documenteren — externe links breken na migratie.",
             "M365 Engineer"),
            ("Verouderde data identificeren",
             f"Met {company} afstemmen welke data gearchiveerd of weggegooid kan worden vóór de migratie.",
             "Projectleider"),
        ]),
        ("INVENTORY — TEAMS", [
            ("Teams & kanalen exporteren",
             "Alle teams, kanalen (incl. private/shared channels!), leden en eigenaren exporteren. Private channels hebben eigen SharePoint-sites.",
             "M365 Engineer"),
            ("Chat-historie beoordelen",
             "Bepalen of 1:1/groepschats gemigreerd moeten worden (duur en beperkt mogelijk; vaak alleen teamkanaal-berichten migreren).",
             "Projectleider"),
            ("Apps, tabs & connectoren",
             "Geïnstalleerde Teams-apps, tabs (Planner, OneNote, websites) en webhooks per team documenteren — moeten handmatig opnieuw.",
             "M365 Engineer"),
            ("Telefonie/voice (indien aanwezig)",
             "Teams Phone-nummers, call queues en auto attendants inventariseren; nummerporting plannen.",
             "M365 Engineer"),
        ]),
        ("INVENTORY — APPLICATIES & OVERIG", [
            ("Enterprise apps & SSO",
             "Applicaties die op de bron-tenant inloggen via Entra SSO (SaaS-apps, VPN, etc.) — moeten opnieuw gekoppeld worden aan Kelso.",
             "M365 Engineer"),
            ("Power Platform",
             "Power Automate flows, Power Apps en Power BI-rapporten/werkruimtes inventariseren — migreren zelden automatisch mee.",
             "M365 Engineer"),
            ("Intune & compliance policies",
             "Device-configuratie, compliance policies, conditional access en app protection policies documenteren ter herbouw in Kelso.",
             "M365 Engineer"),
            ("SMTP-relay & multifunctionals",
             "Printers, scanners en applicaties die mailen via de bron-tenant — herconfigureren naar Kelso.",
             "Werkplekbeheer"),
            ("Risico-analyse afronden",
             "Alle bevindingen samenvatten op het tabblad 'Risico's & Issues' en bespreken met de klant.",
             "Projectleider"),
        ]),
    ]


def phases_identiteit(company):
    return [
        ("GEBRUIKERS AANMAKEN IN KELSO", [
            ("User-mapping tabel opstellen",
             f"Bron-UPN → doel-UPN mapping voor alle gebruikers (tabblad 'User Mapping'). Dit is de basis voor de migratietool.",
             "M365 Engineer"),
            ("Gebruikers provisionen in Kelso",
             "Accounts aanmaken in de Kelso-tenant (CSV-import, Graph of scripted). Tijdelijk @kelso-onmicrosoft of doel-domein als UPN.",
             "M365 Engineer"),
            ("Licenties toewijzen",
             "Juiste licenties per gebruiker toewijzen zodat mailbox/OneDrive geprovisioned wordt (OneDrive pre-provisioning via PowerShell versnellen).",
             "M365 Engineer"),
            ("Groepen & DL's aanmaken",
             "Distributielijsten, security groups en M365-groepen nabouwen in Kelso met juiste leden/eigenaren.",
             "M365 Engineer"),
            ("Shared & resource mailboxen aanmaken",
             "Gedeelde mailboxen en room/equipment mailboxen aanmaken met dezelfde delegaties als in de bron.",
             "M365 Engineer"),
        ]),
        ("BEVEILIGING & BELEID", [
            ("MFA/Conditional Access voorbereiden",
             "CA-policies in Kelso controleren; zorgen dat migratie-serviceaccounts uitgezonderd zijn tijdens de migratie.",
             "M365 Engineer"),
            ("Wachtwoordbeleid & SSPR",
             "Initiële wachtwoorden uitgeven via veilig kanaal; Self-Service Password Reset instellen.",
             "M365 Engineer"),
            ("Admin-rollen toekennen",
             "Benodigde beheerrollen in Kelso toewijzen volgens least-privilege.",
             "M365 Engineer"),
        ]),
    ]


def phases_exchange(company):
    return [
        ("PRE-STAGE (VÓÓR CUTOVER)", [
            ("Migratietool configureren",
             f"Bron ({company}) en doel ({TARGET}) endpoints koppelen in de migratietool; mapping-tabel importeren; testbatch met 2-3 mailboxen draaien.",
             "M365 Engineer"),
            ("Pre-stage sync mailboxen",
             "Eerste volledige sync van alle mailboxdata (mail, agenda, contacten, taken) terwijl gebruikers nog op de bron werken. Grote mailboxen eerst starten.",
             "M365 Engineer"),
            ("Archieven migreren",
             "In-Place Archives apart migreren naar archieven in Kelso (controleer tool-ondersteuning).",
             "M365 Engineer"),
            ("Delta-syncs draaien",
             "Dagelijkse delta-syncs tot aan cutover zodat de delta op cutover-dag minimaal is.",
             "M365 Engineer"),
            ("Foutrapportage controleren",
             "Per batch errors/skipped items beoordelen en oplossen (corrupte items, te grote items >150MB, throttling).",
             "M365 Engineer"),
        ]),
        ("CUTOVER-DAG", [
            ("Mailflow bevriezen / final delta",
             "Laatste delta-sync draaien. Optioneel: bron-mailboxen op verzenden blokkeren tijdens het cutover-venster.",
             "M365 Engineer"),
            ("MX-records omzetten",
             f"MX, Autodiscover, SPF, DKIM en DMARC omzetten naar de {TARGET}-tenant. DKIM in Kelso vooraf klaarzetten en activeren.",
             "M365 Engineer"),
            ("Domein verhuizen (indien van toepassing)",
             f"Maildomein losmaken uit {company}-tenant (alle aliassen/UPN's eerst omzetten naar onmicrosoft) en toevoegen + verifiëren in {TARGET}; daarna UPN's en primaire SMTP omzetten.",
             "M365 Engineer"),
            ("Mailflow testen",
             "Inbound/outbound mail testen (intern, extern, naar shared mailboxen, distributielijsten). Mail-tips en handtekeningen controleren.",
             "M365 Engineer"),
            ("Forwarding bron → doel (vangnet)",
             "Tijdelijke forwarding instellen op bron-mailboxen voor mail die nog via oude route binnenkomt.",
             "M365 Engineer"),
        ]),
        ("NA CUTOVER", [
            ("Post-cutover delta",
             "Laatste delta draaien om mail te vangen die tijdens cutover nog in de bron is bezorgd.",
             "M365 Engineer"),
            ("Outlook-profielen vernieuwen",
             "Nieuwe Outlook-profielen op alle werkplekken; cache opnieuw opbouwen; handtekeningen terugzetten.",
             "Servicedesk"),
            ("Mobiele mail herconfigureren",
             "Outlook mobile opnieuw koppelen aan het Kelso-account.",
             "Servicedesk"),
        ]),
    ]


def phases_files(company):
    return [
        ("ONEDRIVE", [
            ("OneDrive pre-provisioning",
             "OneDrives in Kelso vooraf aanmaken (Request-SPOPersonalSite) zodat de migratietool direct kan schrijven.",
             "M365 Engineer"),
            ("OneDrive pre-stage sync",
             "Volledige eerste sync van alle OneDrive-data; daarna delta-syncs tot cutover.",
             "M365 Engineer"),
            ("Versiegeschiedenis & metadata",
             "Tool-instellingen controleren: versies, auteurs en timestamps meenemen waar mogelijk.",
             "M365 Engineer"),
            ("Gedeelde links communiceren",
             "Gebruikers informeren dat bestaande deellinks breken en opnieuw gedeeld moet worden.",
             "Servicedesk"),
            ("OneDrive-client herkoppelen",
             "Na cutover op elke werkplek OneDrive ontkoppelen van bron en koppelen aan Kelso-account (sync-conflicten vermijden: oude map hernoemen).",
             "Werkplekbeheer"),
        ]),
        ("SHAREPOINT", [
            ("Site-structuur aanmaken in Kelso",
             "Doelsites aanmaken (of door tool laten aanmaken) volgens afgestemde structuur; site-eigenaren instellen.",
             "M365 Engineer"),
            ("SharePoint pre-stage sync",
             "Volledige sync van alle te migreren sites incl. documentbibliotheken, lijsten en permissies; delta's tot cutover.",
             "M365 Engineer"),
            ("Permissies valideren",
             "Steekproef per site: kloppen de rechten (eigenaren/leden/bezoekers, unieke permissies)?",
             "M365 Engineer"),
            ("Snelkoppelingen & gesyncte bibliotheken",
             "Gebruikers met gesyncte SharePoint-bibliotheken: oude sync verwijderen, nieuwe sync naar Kelso-sites instellen.",
             "Werkplekbeheer"),
            ("Final delta & verificatie",
             "Laatste delta op cutover; itemcounts bron vs. doel vergelijken; afwijkingen rapporteren.",
             "M365 Engineer"),
        ]),
    ]


def phases_teams(company):
    return [
        ("TEAMS — STRUCTUUR & DATA", [
            ("Teams aanmaken in Kelso",
             "Teams + kanalen (incl. private/shared channels) nabouwen of door de migratietool laten aanmaken; eigenaren en leden koppelen via de mapping-tabel.",
             "M365 Engineer"),
            ("Kanaalbestanden migreren",
             "Bestanden van elk kanaal (SharePoint achter het team) migreren; private channel-sites niet vergeten.",
             "M365 Engineer"),
            ("Kanaalberichten migreren",
             "Kanaalconversaties migreren indien gewenst (vaak als HTML-archief of via tool met beperkingen). Verwachtingen bij gebruikers managen.",
             "M365 Engineer"),
            ("Chats (1:1/groep)",
             "Indien in scope: chat-migratie draaien (beperkt; alternatief = leesbaar archief). Beslissing documenteren.",
             "Projectleider"),
            ("Tabs, apps & koppelingen herstellen",
             "Planner, OneNote, websites-tabs en app-koppelingen per team handmatig terugzetten aan de hand van de inventory.",
             "M365 Engineer"),
        ]),
        ("TEAMS — NAZORG", [
            ("Teams-client resetten",
             "Gebruikers uitloggen/cache wissen en aanmelden met Kelso-account; vergaderingen in agenda's opnieuw aanmaken indien Teams-links naar bron-tenant wijzen.",
             "Servicedesk"),
            ("Terugkerende vergaderingen",
             "Terugkerende Teams-meetings bevatten links naar de oude tenant — organisatoren laten her-aanmaken na cutover.",
             "Servicedesk"),
            ("Telefonie omzetten (indien aanwezig)",
             "Nummers porteren/verhuizen, call queues en auto attendants nabouwen in Kelso.",
             "M365 Engineer"),
        ]),
    ]


def phases_cutover(company):
    return [
        ("CUTOVER-CHECKLIST (GO/NO-GO)", [
            ("Go/No-Go beslissing",
             "Checklist doorlopen met projectteam: pre-stage syncs ≥95% compleet, geen blokkerende fouten, communicatie verstuurd, support klaar.",
             "Projectleider"),
            ("Eindgebruikerscommunicatie versturen",
             "Laatste instructiemail: tijdstip, wat te doen maandagochtend, waar hulp te krijgen.",
             "Projectleider"),
            ("Final delta alle workloads",
             "Laatste delta-sync mail, OneDrive, SharePoint en Teams.",
             "M365 Engineer"),
            ("DNS-cutover uitvoeren",
             "MX/Autodiscover/SPF/DKIM/DMARC omzetten (zie Exchange-tabblad).",
             "M365 Engineer"),
            ("UPN & primaire SMTP omzetten",
             "Gebruikers in Kelso omzetten naar het productiedomein als primair adres.",
             "M365 Engineer"),
            ("Bron-tenant beperken",
             "Sign-in blokkeren op bron-accounts (na validatie) om verwarring en dubbel gebruik te voorkomen.",
             "M365 Engineer"),
        ]),
        ("VALIDATIE NA CUTOVER", [
            ("Mailflow-validatie",
             "Test intern/extern/inbound/outbound, shared mailboxen, DL's, handtekeningen, agenda-delegaties.",
             "M365 Engineer"),
            ("Data-steekproeven",
             "Per gebruikersgroep steekproef op mail, OneDrive, SharePoint en Teams-bestanden (counts en recente items).",
             "M365 Engineer"),
            ("Werkplekken omgezet",
             "Alle pc's: nieuw profiel, OneDrive gekoppeld, Teams aangemeld, printers/scan-to-mail werkt.",
             "Werkplekbeheer"),
            ("Hypercare-periode starten",
             "1-2 weken verhoogde support; issues loggen op het tabblad 'Risico's & Issues'.",
             "Servicedesk"),
        ]),
        ("AFRONDING & DECOMMISSIE", [
            ("Migratierapport opleveren",
             f"Eindrapport: gemigreerde aantallen, openstaande punten, afwijkingen, lessons learned voor {company}.",
             "Projectleider"),
            ("Forwarding & service-accounts opruimen",
             "Tijdelijke forwarding, migratie-accounts en app-registraties verwijderen in beide tenants.",
             "M365 Engineer"),
            ("Bron-tenant afbouwen",
             "Na afgesproken bewaarperiode: licenties opzeggen, data-export/back-up veiligstellen, tenant opheffen of slapend maken.",
             "M365 Engineer"),
            ("Documentatie & overdracht",
             "Beheerdocumentatie Kelso-omgeving bijwerken en overdragen aan beheer.",
             "Projectleider"),
        ]),
    ]


# ---------- Speciale tabbladen ----------

def overview_sheet(wb, company):
    ws = wb.create_sheet("Overzicht", 0)
    ws.sheet_properties.tabColor = C_DARK
    banner(ws, company, "Projectoverzicht & voortgang per fase", 6)

    info = [
        ("Bron-tenant", company),
        ("Doel-tenant", TARGET),
        ("Projectleider", ""),
        ("M365 Engineer", ""),
        ("Migratietool", ""),
        ("Geplande cutover-datum", ""),
        ("Aantal gebruikers", ""),
        ("Aantal mailboxen (incl. shared)", ""),
        ("Totale datagrootte (GB)", ""),
        ("Status project", "Niet gestart"),
    ]
    row = 4
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    c = ws.cell(row=row, column=1, value="PROJECTGEGEVENS")
    c.font = F_PHASE; c.fill = FILL_ACCENT
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(1, 7):
        ws.cell(row=row, column=col).fill = FILL_ACCENT
    row += 1
    for label, val in info:
        lc = ws.cell(row=row, column=1, value=label)
        lc.font = F_BOLD; lc.fill = FILL_LIGHT; lc.border = BORDER
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        vc = ws.cell(row=row, column=2, value=val)
        vc.font = F_BODY; vc.border = BORDER
        ws.cell(row=row, column=3).border = BORDER
        row += 1

    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    c = ws.cell(row=row, column=1, value="VOORTGANG PER FASE")
    c.font = F_PHASE; c.fill = FILL_ACCENT
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(1, 7):
        ws.cell(row=row, column=col).fill = FILL_ACCENT
    row += 1
    hdrs = ["Fase", "Tabblad", "Status", "Startdatum", "Einddatum", "Opmerkingen"]
    for i, h in enumerate(hdrs, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = F_HDR; c.fill = FILL_MED; c.alignment = CENTER; c.border = BORDER
    row += 1
    first_status = row
    fases = [
        ("1. Voorbereiding", "1. Voorbereiding"),
        ("2. Analyse & Inventory", "2. Analyse & Inventory"),
        ("3. Identiteit & Gebruikers", "3. Identiteit"),
        ("4. Exchange Online", "4. Exchange"),
        ("5. OneDrive & SharePoint", "5. OneDrive-SharePoint"),
        ("6. Teams", "6. Teams"),
        ("7. Cutover & Nazorg", "7. Cutover & Nazorg"),
    ]
    for naam, tab in fases:
        ws.cell(row=row, column=1, value=naam).font = F_BOLD
        ws.cell(row=row, column=2, value=tab).font = F_BODY
        ws.cell(row=row, column=3, value="Niet gestart").alignment = CENTER
        for col in range(1, 7):
            ws.cell(row=row, column=col).border = BORDER
        row += 1
    status_validation(ws, "C", first_status, row - 1)

    widths = [32, 26, 16, 14, 14, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def mapping_sheet(wb, company):
    ws = wb.create_sheet("User Mapping")
    ws.sheet_properties.tabColor = "7030A0"
    headers = ["Nr", "Displaynaam", f"Bron UPN ({company})", f"Doel UPN ({TARGET})",
               "Licentie doel", "Mailbox type", "Mailboxgrootte (GB)", "OneDrive (GB)",
               "Gemigreerd?", "Notities"]
    widths = [6, 26, 34, 34, 18, 16, 16, 14, 13, 30]
    banner(ws, company, "Gebruikersmapping bron → doel (vul aan vanuit de inventory-export)", len(headers))
    header_row(ws, 4, headers, widths)
    for r in range(5, 55):
        ws.cell(row=r, column=1, value=r - 4).alignment = CENTER
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).border = BORDER
            ws.cell(row=r, column=col).font = F_BODY
    dv = DataValidation(type="list", formula1='"Ja,Nee,Bezig"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("I5:I54")
    dv2 = DataValidation(type="list", formula1='"User,Shared,Room,Equipment"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add("F5:F54")
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:J54"
    return ws


def risk_sheet(wb, company):
    ws = wb.create_sheet("Risico's & Issues")
    ws.sheet_properties.tabColor = "C00000"
    headers = ["Nr", "Datum", "Type", "Omschrijving", "Impact", "Prioriteit",
               "Eigenaar", "Status", "Oplossing / Mitigatie"]
    widths = [6, 12, 12, 50, 30, 11, 18, 13, 45]
    banner(ws, company, "Risico- en issue-log — houd hier alle bevindingen en incidenten bij", len(headers))
    header_row(ws, 4, headers, widths)
    seed = [
        ("Risico", "Externe deellinks (SharePoint/OneDrive) werken niet meer na migratie.",
         "Gebruikers/externen verliezen toegang tot gedeelde bestanden.",
         "Communiceren; belangrijkste shares na migratie opnieuw delen."),
        ("Risico", "Terugkerende Teams-vergaderlinks verwijzen naar de oude tenant.",
         "Vergaderingen niet toegankelijk na cutover.",
         "Organisatoren laten her-aanmaken na cutover."),
        ("Risico", "Throttling door Microsoft tijdens grote datamigraties.",
         "Migratie duurt langer dan gepland.",
         "Vroeg starten met pre-stage; migratiegolven spreiden."),
        ("Risico", "Items >150 MB of corrupte items worden overgeslagen.",
         "Ontbrekende bestanden/mail items.",
         "Foutrapporten per batch controleren; handmatig nabehandelen."),
        ("Risico", "Domein kan niet verhuizen door achtergebleven objecten.",
         "Cutover-vertraging.",
         "Vooraf alle aliassen/UPN's omzetten naar .onmicrosoft.com en valideren."),
    ]
    row = 5
    for i, (typ, oms, imp, mit) in enumerate(seed, 1):
        vals = [i, "", typ, oms, imp, "Hoog", "", "Open", mit]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = F_BODY; c.border = BORDER
            c.alignment = WRAP if col in (4, 5, 9) else CENTER
        row += 1
    for r in range(row, row + 20):
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).border = BORDER
    last = row + 19
    dv = DataValidation(type="list", formula1='"Open,Bezig,Opgelost,Geaccepteerd"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"H5:H{last}")
    dvt = DataValidation(type="list", formula1='"Risico,Issue,Beslissing,Actie"', allow_blank=True)
    ws.add_data_validation(dvt)
    dvt.add(f"C5:C{last}")
    prio_validation(ws, "F", 5, last)
    ws.conditional_formatting.add(
        f"H5:H{last}", CellIsRule(operator="equal", formula=['"Opgelost"'],
                                  fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(
        f"H5:H{last}", CellIsRule(operator="equal", formula=['"Open"'],
                                  fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:I{last}"
    return ws


def build_workbook(company):
    wb = Workbook()
    wb.remove(wb.active)
    overview_sheet(wb, company)
    task_sheet(wb, company, "1. Voorbereiding", "4472C4",
               "Fase 1 — Voorbereiding: governance, toegang, licenties, DNS en werkplek",
               phases_voorbereiding(company))
    task_sheet(wb, company, "2. Analyse & Inventory", "70AD47",
               "Fase 2 — Analyse & Inventory: breng de volledige bron-tenant in kaart",
               phases_inventory(company))
    mapping_sheet(wb, company)
    task_sheet(wb, company, "3. Identiteit", "FFC000",
               "Fase 3 — Identiteit & gebruikers provisionen in de Kelso-tenant",
               phases_identiteit(company))
    task_sheet(wb, company, "4. Exchange", "ED7D31",
               "Fase 4 — Exchange Online migratie: pre-stage, cutover en nazorg",
               phases_exchange(company))
    task_sheet(wb, company, "5. OneDrive-SharePoint", "5B9BD5",
               "Fase 5 — OneDrive & SharePoint datamigratie",
               phases_files(company))
    task_sheet(wb, company, "6. Teams", "7030A0",
               "Fase 6 — Microsoft Teams migratie",
               phases_teams(company))
    task_sheet(wb, company, "7. Cutover & Nazorg", "C00000",
               "Fase 7 — Cutover, validatie, hypercare en decommissie van de bron-tenant",
               phases_cutover(company))
    risk_sheet(wb, company)
    return wb


def safe_name(name):
    for ch in '\\/:*?"<>|&':
        name = name.replace(ch, "-")
    return name.replace(" ", "_")


if __name__ == "__main__":
    os.makedirs(OUT_DIR, exist_ok=True)
    for comp in COMPANIES:
        wb = build_workbook(comp)
        path = os.path.join(OUT_DIR, f"M365_Migratie_{safe_name(comp)}_naar_{TARGET}.xlsx")
        wb.save(path)
        print("OK", path)
    print("Klaar:", len(COMPANIES), "werkboeken in", OUT_DIR)
